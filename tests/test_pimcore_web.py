import csv
import io
import json
from unittest.mock import Mock, patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from picorgftp_sql import web_data
from picorgftp_sql.services.pimcore_service import PimcoreApiError, PimcoreConflictError
from picorgftp_sql.sqlite_store import SqliteStore
from picorgftp_sql.web import app as web_app


def test_settings_snapshot_hides_pimcore_api_key():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["enabled"] = True
    cfg["pimcore"]["api_key"] = "secret"
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "load_users", return_value=[]),
    ):
        snapshot = web_data.settings_snapshot()

    assert snapshot["pimcore"]["api_key_set"] is True
    assert "api_key" not in snapshot["pimcore"]


def test_update_settings_preserves_blank_pimcore_api_key_and_saves_mapping():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["api_key"] = "saved-secret"
    saved = []
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "save_config",
            side_effect=lambda payload, **kwargs: saved.append(
                json.loads(json.dumps(payload))
            ),
        ),
        patch.object(web_data.config, "initialize_config", return_value=cfg),
        patch.object(web_data, "settings_snapshot", return_value={}),
    ):
        web_data.update_settings(
            {
                "pimcore": {
                    "enabled": True,
                    "api_key": "",
                    "base_url": "http://10.10.0.5",
                    "field_mappings": [
                        {
                            "source": "EAN",
                            "pimcore_field": "EAN",
                            "type": "input",
                            "required": True,
                            "parser": "text",
                        }
                    ],
                }
            }
        )

    assert saved[0]["pimcore"]["api_key"] == "saved-secret"
    assert saved[0]["pimcore"]["field_mappings"][0]["source"] == "EAN"


def test_parse_csv_headers_supports_semicolon_and_quoted_labels():
    raw = (
        b'SKU;EAN;"TOTAL WEIGHT";"TOTAL VOLUME [m2]"\r\n'
        b"ABC;5904804578169;62,5;1,2\r\n"
    )

    assert web_data.parse_pimcore_csv_headers(raw) == [
        "SKU",
        "EAN",
        "TOTAL WEIGHT",
        "TOTAL VOLUME [m2]",
    ]


def test_pimcore_settings_test_route_returns_structured_report():
    client = TestClient(web_app.app)
    report = {
        "ok": False,
        "checks": [{"key": "mapping_fields", "status": "error"}],
        "total_ms": 4,
    }
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(web_app, "test_pimcore_settings", return_value=report),
    ):
        response = client.post("/api/settings/pimcore/test")

    assert response.status_code == 200
    assert response.json() == report


def test_admin_can_test_sql_profile_route():
    client = TestClient(web_app.app)
    expected = {"ok": True, "message": "Polaczenie SQL dziala."}
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(
            web_app,
            "test_sql_profile_connection",
            return_value=expected,
        ) as test_profile,
    ):
        response = client.post("/api/settings/sql-profiles/stock/test")

    assert response.status_code == 200
    assert response.json() == expected
    test_profile.assert_called_once_with("stock")


def test_settings_diagnostic_persists_full_detail_but_returns_public_report():
    report = {
        "ok": False,
        "checks": [
            {
                "key": "server_info",
                "status": "error",
                "response_excerpt": "short trace",
                "response_detail": "complete sanitized trace",
            }
        ],
    }
    with (
        patch.object(web_data, "run_settings_test", return_value=report),
        patch.object(web_data, "record_history") as record,
    ):
        result = web_data.test_pimcore_settings({}, "admin")

    assert "response_detail" not in result["checks"][0]
    persisted = record.call_args.kwargs["details"]["pimcore_settings_test"]
    assert persisted["checks"][0]["response_detail"] == "complete sanitized trace"


def test_discovery_uses_unsaved_key_without_persisting_or_returning_it():
    captured = {}
    fake_client = Mock()
    with (
        patch.object(web_data.config, "CONFIG", {"pimcore": {"api_key": "saved"}}),
        patch.object(web_data, "PimcoreClient", return_value=fake_client) as client_type,
        patch.object(
            web_data,
            "discover_classes",
            return_value=[{"id": "7", "name": "product"}],
        ) as discover,
    ):
        client_type.side_effect = lambda settings: (
            captured.setdefault("settings", settings),
            fake_client,
        )[1]
        result = web_data.discover_pimcore_classes(
            {"base_url": "http://10.10.0.5", "api_key": "temporary"}
        )

    assert captured["settings"]["api_key"] == "temporary"
    assert result == {"items": [{"id": "7", "name": "product"}]}
    assert "temporary" not in json.dumps(result)
    discover.assert_called_once_with(fake_client)


def test_complete_setup_saves_only_after_successful_report():
    payload = {
        "base_url": "http://10.10.0.5",
        "api_key": "secret",
        "class_id": "7",
        "class_name": "product",
        "parent_id": "6626",
        "parent_path": "/Produkty",
        "field_mappings": [
            {
                "source": "EAN",
                "pimcore_field": "EAN",
                "type": "input",
                "required": True,
                "parser": "text",
            }
        ],
    }
    with (
        patch.object(web_data, "test_pimcore_settings", return_value={"ok": True, "checks": []}),
        patch.object(
            web_data,
            "update_settings",
            return_value={"pimcore": {"setup_complete": True}},
        ) as save,
    ):
        result = web_data.complete_pimcore_setup(payload, "admin")

    assert result["saved"] is True
    saved = save.call_args.args[0]["pimcore"]
    assert saved["setup_complete"] is True
    assert saved["enabled"] is True
    assert saved["object_key_template"] == "{EAN}"


def test_complete_setup_does_not_save_after_failed_report():
    with (
        patch.object(web_data, "test_pimcore_settings", return_value={"ok": False, "checks": []}),
        patch.object(web_data, "update_settings") as save,
    ):
        result = web_data.complete_pimcore_setup({"api_key": "secret"}, "admin")

    assert result == {"saved": False, "report": {"ok": False, "checks": []}}
    save.assert_not_called()


def test_pimcore_discovery_and_setup_routes_are_admin_only():
    client = TestClient(web_app.app)
    admin = {"username": "admin", "role": "admin"}
    with (
        patch.object(web_app, "_require_admin", return_value=admin) as require_admin,
        patch.object(
            web_app,
            "discover_pimcore_classes",
            return_value={"items": [{"id": "7", "name": "product"}]},
        ),
        patch.object(web_app, "discover_pimcore_fields", return_value={"items": []}),
        patch.object(web_app, "discover_pimcore_folders", return_value={"items": []}),
        patch.object(
            web_app,
            "complete_pimcore_setup",
            return_value={"saved": True, "report": {"ok": True}},
        ),
    ):
        classes = client.post("/api/settings/pimcore/discover/classes", json={"settings": {}})
        fields = client.post(
            "/api/settings/pimcore/discover/fields",
            json={"settings": {}, "class_id": "7"},
        )
        folders = client.post("/api/settings/pimcore/discover/folders", json={"settings": {}})
        saved = client.post("/api/settings/pimcore/setup", json={"settings": {}})

    assert classes.status_code == fields.status_code == folders.status_code == saved.status_code == 200
    assert require_admin.call_count == 4


def test_folder_discovery_route_degrades_to_empty_list_on_pimcore_error():
    client = TestClient(web_app.app)
    error = PimcoreApiError(
        "Pimcore zwrocil HTTP 502.",
        "/webservice/rest/object-list",
        status_code=502,
    )
    with (
        patch.object(web_app, "_require_admin", return_value={"username": "admin", "role": "admin"}),
        patch.object(web_app, "discover_pimcore_folders", side_effect=error),
    ):
        response = client.post("/api/settings/pimcore/discover/folders", json={"settings": {}})

    assert response.status_code == 200
    assert response.json()["items"] == []
    assert response.json()["warning"]["status_code"] == 502


def test_pimcore_csv_headers_route_parses_uploaded_file():
    client = TestClient(web_app.app)
    with patch.object(
        web_app,
        "_require_admin",
        return_value={"username": "admin", "role": "admin"},
    ):
        response = client.post(
            "/api/settings/pimcore/import-csv-headers",
            files={
                "file": (
                    "products.csv",
                    b"SKU;EAN\r\nABC;5904804578169\r\n",
                    "text/csv",
                )
            },
        )

    assert response.status_code == 200
    assert response.json() == {"headers": ["SKU", "EAN"]}


def test_pimcore_operation_history_reads_persisted_audit_records():
    operation = {
        "operation_id": "op-1",
        "operation_type": "test",
        "username": "admin",
        "status": "partial",
        "started_at": 15.0,
        "events": [{"sequence": 1, "stage": "delete", "message": "HTTP 403"}],
    }
    records = [
        {"action": "entry_save", "details": {}},
        {"action": "pimcore_test_create", "details": {"pimcore_operation": operation}},
    ]
    with patch.object(web_data, "_load_history_records", return_value=records):
        result = web_data.pimcore_operation_history(
            operation_type="test",
            result="partial",
            user="admin",
            query="HTTP 403",
            date_from=10,
            date_to=20,
        )

    assert result == {"items": [operation], "count": 1}


def test_pimcore_test_run_routes_forward_admin_and_sequence():
    client = TestClient(web_app.app)
    user = {"username": "admin", "role": "admin"}
    with (
        patch.object(web_app, "_require_admin", return_value=user),
        patch.object(
            web_app,
            "start_pimcore_test_create",
            return_value={"operation_id": "op-1", "status": "queued"},
        ),
        patch.object(
            web_app,
            "pimcore_operation_status",
            return_value={"operation_id": "op-1", "events": [], "status": "running"},
        ) as status,
    ):
        started = client.post(
            "/api/settings/pimcore/test-create-runs",
            json={"values": {"EAN": "5904804578169"}, "cleanup_policy": "delete"},
        )
        polled = client.get(
            "/api/settings/pimcore/test-create-runs/op-1?after_sequence=4"
        )

    assert started.status_code == 200
    assert started.json()["operation"]["operation_id"] == "op-1"
    assert polled.status_code == 200
    status.assert_called_once_with("op-1", 4)


def test_missing_pimcore_operation_returns_404():
    client = TestClient(web_app.app)
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(web_app, "pimcore_operation_status", return_value=None),
    ):
        response = client.get("/api/settings/pimcore/test-create-runs/missing")

    assert response.status_code == 404


def test_pimcore_history_route_forwards_all_filters():
    client = TestClient(web_app.app)
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(
            web_app,
            "pimcore_operation_history",
            return_value={"items": [], "count": 0},
        ) as history,
    ):
        response = client.get(
            "/api/settings/pimcore/operations?"
            "operation_type=test&result=partial&user=admin&query=5904&"
            "date_from=10&date_to=20&limit=25"
        )

    assert response.status_code == 200
    history.assert_called_once_with(
        operation_type="test",
        result="partial",
        user="admin",
        query="5904",
        date_from=10.0,
        date_to=20.0,
        limit=25,
    )


def test_admin_can_export_pimcore_submissions_as_json():
    client = TestClient(web_app.app)
    expected = {"items": [{"operation_id": "op-1"}], "format": "json"}
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(
            web_app,
            "export_pimcore_submissions",
            return_value=expected,
        ) as export,
    ):
        response = client.get(
            "/api/settings/pimcore/submissions/export?format=json&user=operator"
        )

    assert response.status_code == 200
    assert response.json() == expected
    export.assert_called_once()


def test_export_pimcore_submissions_as_csv_uses_pimcore_field_labels_only():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["field_mappings"] = [
        {
            "source": "EAN",
            "label": "Kod EAN",
            "pimcore_field": "ean",
            "type": "input",
            "parser": "text",
        },
        {
            "source": "STOCK",
            "label": "Stan",
            "pimcore_field": "stock",
            "type": "numeric",
            "parser": "integer",
        },
    ]
    store = Mock()
    store.query_pimcore_submissions.return_value = [
        {
            "operation_id": "op-1",
            "operation_type": "manual_create",
            "username": "operator",
            "ean": "5901234567890",
            "status": "completed",
            "values": {"EAN": "5901234567890", "STOCK": "12", "UNMAPPED": "hidden"},
            "payload": {"className": "Product"},
            "result": {"object_id": 91},
            "warnings": [{"message": "hidden"}],
            "created_at": "2026-07-06T12:00:00.000Z",
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "_active_sqlite_store", return_value=store),
    ):
        exported = web_data.export_pimcore_submissions(export_format="csv")

    assert exported["format"] == "csv"
    rows = list(csv.reader(io.StringIO(exported["content"])))
    assert rows == [["Kod EAN", "Stan"], ["5901234567890", "12"]]
    assert "operation_id" not in exported["content"]
    assert "payload" not in exported["content"]
    assert "UNMAPPED" not in exported["content"]


def test_export_pimcore_submissions_as_xlsx_uses_pimcore_field_labels_only():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["field_mappings"] = [
        {
            "source": "EAN",
            "label": "Kod EAN",
            "pimcore_field": "ean",
            "type": "input",
            "parser": "text",
        },
        {
            "source": "STOCK",
            "label": "Stan",
            "pimcore_field": "stock",
            "type": "numeric",
            "parser": "integer",
        },
    ]
    store = Mock()
    store.query_pimcore_submissions.return_value = [
        {
            "operation_id": "op-1",
            "operation_type": "manual_create",
            "username": "operator",
            "ean": "5901234567890",
            "status": "completed",
            "values": {"STOCK": "12"},
            "payload": {"className": "Product"},
            "result": {"object_id": 91},
            "warnings": [],
            "created_at": "2026-07-06T12:00:00.000Z",
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "_active_sqlite_store", return_value=store),
    ):
        exported = web_data.export_pimcore_submissions(export_format="xlsx")

    assert exported["format"] == "xlsx"
    workbook = load_workbook(io.BytesIO(exported["content"]))
    sheet = workbook.active
    assert [cell.value for cell in sheet[1]] == ["Kod EAN", "Stan"]
    assert [cell.value for cell in sheet[2]] == ["5901234567890", "12"]


def test_export_pimcore_submissions_ignores_technical_blocks():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["field_mappings"] = [
        {
            "source": "EAN",
            "label": "Kod EAN",
            "pimcore_field": "ean",
            "type": "input",
            "parser": "text",
        },
        {
            "source": "STOCK",
            "label": "Stan",
            "pimcore_field": "stock",
            "type": "numeric",
            "parser": "integer",
        },
    ]
    store = Mock()
    store.query_pimcore_submissions.return_value = [
        {
            "operation_id": "op-1",
            "operation_type": "manual_create",
            "username": "operator",
            "ean": "5901234567890",
            "status": "completed",
            "values": {"EAN": "5901234567890", "STOCK": "12"},
            "payload": {
                "className": "Product",
                "elements": [{"name": "EAN", "value": "5901234567890"}],
            },
            "result": {"object_id": 91, "object": {"path": "/Produkty/5901234567890"}},
            "warnings": [{"code": "missing_sql", "message": "Brak danych SQL"}],
            "created_at": "2026-07-06T12:00:00.000Z",
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "_active_sqlite_store", return_value=store),
    ):
        exported = web_data.export_pimcore_submissions(export_format="csv")

    rows = list(csv.reader(io.StringIO(exported["content"])))
    assert rows == [["Kod EAN", "Stan"], ["5901234567890", "12"]]
    assert all("payload." not in column for column in rows[0])
    assert all("result." not in column for column in rows[0])
    assert all("warnings" not in column for column in rows[0])


def test_admin_can_export_pimcore_submissions_as_xlsx_response():
    client = TestClient(web_app.app)
    with (
        patch.object(
            web_app,
            "_require_admin",
            return_value={"username": "admin", "role": "admin"},
        ),
        patch.object(
            web_app,
            "export_pimcore_submissions",
            return_value={"format": "xlsx", "content": b"xlsx-bytes", "count": 1},
        ) as export,
    ):
        response = client.get("/api/settings/pimcore/submissions/export?format=xlsx")

    assert response.status_code == 200
    assert response.content == b"xlsx-bytes"
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "pimcore-submissions.xlsx" in response.headers["content-disposition"]
    export.assert_called_once()


def test_product_status_returns_disabled_without_network_call():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"]["enabled"] = False

    with patch.object(web_data.config, "CONFIG", cfg):
        assert web_data.find_pimcore_product_by_ean("5904804578169") == {
            "enabled": False,
            "setup_complete": False,
            "exists": False,
            "object": None,
            "form_schema": [],
        }


def test_runtime_form_schema_includes_sql_mapping_metadata():
    settings_payload = web_data.normalize_pimcore_settings(
        {
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stan",
                    "pimcore_field": "stock",
                    "type": "input",
                    "parser": "text",
                    "value_template": "SQL",
                    "sql_query": "SELECT qty FROM stock WHERE ean = {ean}",
                    "sql_profile_id": "stock-db",
                }
            ]
        }
    )

    schema = web_data._pimcore_runtime_form_schema(settings_payload)

    assert schema[0]["value_template"] == "SQL"
    assert schema[0]["sql_query"] == "SELECT qty FROM stock WHERE ean = {ean}"
    assert schema[0]["sql_profile_id"] == "stock-db"


def test_runtime_form_schema_includes_layout_and_display_order():
    settings_payload = web_data.normalize_pimcore_settings(
        {
            "field_mappings": [
                {
                    "source": "DESCRIPTION",
                    "label": "Opis",
                    "pimcore_field": "description",
                    "type": "textarea",
                    "parser": "text",
                    "layout_group": "Opis",
                    "layout_order": 20,
                },
                {
                    "source": "TITLE",
                    "label": "Tytul",
                    "pimcore_field": "title",
                    "type": "input",
                    "parser": "text",
                    "layout_group": "Dane podstawowe",
                    "layout_order": 10,
                },
            ]
        }
    )

    schema = web_data._pimcore_runtime_form_schema(settings_payload)

    assert [item["source"] for item in schema] == ["TITLE", "DESCRIPTION"]
    assert schema[0]["layout_group"] == "Dane podstawowe"
    assert schema[0]["layout_order"] == 10
    assert "layout_width" not in schema[0]
    assert schema[1]["layout_group"] == "Opis"
    assert "layout_width" not in schema[1]


def test_runtime_status_is_disabled_when_setup_is_incomplete():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": False})
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "find_product_by_ean") as lookup,
    ):
        result = web_data.find_pimcore_product_by_ean("5904804578169")

    assert result == {
        "enabled": False,
        "setup_complete": False,
        "exists": False,
        "object": None,
        "form_schema": [],
    }
    lookup.assert_not_called()


def test_bootstrap_exposes_only_runtime_pimcore_flags():
    client = TestClient(web_app.app)
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(
            web_app,
            "_current_user_payload",
            return_value={"username": "operator", "role": "user"},
        ),
        patch.object(web_app, "load_web_data", return_value={}),
        patch.object(
            web_app,
            "pimcore_runtime_capabilities",
            return_value={"enabled": True, "setup_complete": True},
        ),
    ):
        response = client.get("/api/bootstrap")

    assert response.json()["pimcore"] == {"enabled": True, "setup_complete": True}


def test_runtime_create_route_allows_logged_in_user_and_returns_created_object():
    client = TestClient(web_app.app)
    expected = {
        "created": True,
        "duplicate": False,
        "object": {"id": 91, "key": "ABC", "path": "/Produkty/ABC"},
    }

    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(web_app, "create_pimcore_product", return_value=expected) as create,
    ):
        response = client.post(
            "/api/pimcore/products",
            json={"values": {"SKU": "ABC", "EAN": "5904804578169"}},
        )

    assert response.status_code == 200
    assert response.json() == expected
    create.assert_called_once_with({"SKU": "ABC", "EAN": "5904804578169"}, "operator")


def test_runtime_create_route_ignores_browser_supplied_integration_results():
    client = TestClient(web_app.app)
    integrations = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "success",
                "elapsed_ms": 8,
                "warning_codes": [],
                "error": "",
            }
        ]
    }
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(
            web_app,
            "create_pimcore_product",
            return_value={"created": True, "object": {"id": 91}},
        ) as create,
    ):
        response = client.post(
            "/api/pimcore/products",
            json={"values": {"EAN": "5904804578169"}, "integration_results": integrations},
        )

    assert response.status_code == 200
    create.assert_called_once_with({"EAN": "5904804578169"}, "operator")


def test_edit_adapter_requires_enabled_complete_setup():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": False, "setup_complete": True})
    with patch.object(web_data.config, "CONFIG", cfg):
        with pytest.raises(ValueError, match="wylaczona"):
            web_data.get_pimcore_product_for_edit(91)


def test_update_adapter_persists_manual_update_audit():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    expected = {
        "object": {"id": 91, "path": "/Produkty/5904"},
        "values": {"EAN": "5904804578169"},
    }
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "update_product", return_value=expected),
        patch.object(web_data, "_persist_pimcore_operation") as persist,
    ):
        result = web_data.update_pimcore_product(
            91,
            "100",
            {"EAN": "5904804578169"},
            "operator",
        )

    assert result == expected
    report = persist.call_args.args[0]
    assert report["operation_type"] == "manual_update"
    assert report["username"] == "operator"


def test_update_adapter_emits_failure_diagnostics_for_manual_update_error():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    error = RuntimeError("Pimcore connection lost")
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "update_product", side_effect=error),
        patch.object(web_data, "_persist_pimcore_operation") as persist,
        patch.object(web_data, "emit_event") as emit_event,
        pytest.raises(RuntimeError, match="Pimcore connection lost"),
    ):
        web_data.update_pimcore_product(
            91,
            "100",
            {"EAN": "5904804578169"},
            "operator",
        )

    assert persist.call_args.args[0]["status"] == "failed"
    event = next(
        call.kwargs
        for call in emit_event.call_args_list
        if call.kwargs["event_type"] == "integration.pimcore.completed"
    )
    assert event["severity"] == "error"
    assert event["exception"] is error
    assert event["recommended_action"] == (
        "Otworz historie operacji Pimcore dla tego EAN i sprawdz zredagowany "
        "zalacznik diagnostyczny."
    )


def test_create_adapter_uses_manual_create_operation_kind():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "create_product",
            return_value={"created": True, "duplicate": False, "object": {"id": 91}},
        ),
        patch.object(web_data, "_persist_pimcore_operation") as persist,
    ):
        web_data.create_pimcore_product({"EAN": "5904804578169"}, "operator")

    assert persist.call_args.args[0]["operation_type"] == "manual_create"


def test_create_adapter_filters_and_attaches_integration_results_to_audit_and_change_set():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stock",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "required": False,
                    "value_template": "SQL",
                    "sql_query": "SELECT 1",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [{"id": "stock", "enabled": True}]
    change_set = {
        "kind": "created",
        "fields": [
            {"key": "EAN", "label": "EAN", "before": None, "after": "5904804578169"}
        ],
    }
    browser_results = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "success",
                "elapsed_ms": 8,
                "warning_codes": ["multiple_rows", {"not": "safe"}],
                "error": "",
                "password": "do-not-persist",
            },
            "not-an-object",
        ],
        "arbitrary": {"token": "do-not-persist"},
    }
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "create_product",
            return_value={
                "created": True,
                "duplicate": False,
                "object": {"id": 91},
                "change_set": change_set,
            },
        ),
        patch.object(web_data, "_persist_pimcore_operation") as persist,
        patch.object(web_data, "_persist_pimcore_submission"),
        patch.object(web_data, "emit_event") as emit_event,
    ):
        result = web_data.create_pimcore_product(
            {"EAN": "5904804578169"},
            "operator",
            browser_results,
        )

    expected = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "success",
                "elapsed_ms": 8,
                "warning_codes": ["multiple_rows"],
                "error": "",
                "required": False,
            }
        ]
    }
    assert result["change_set"]["integrations"] == expected
    report = persist.call_args.args[0]
    assert report["integration_results"] == expected
    assert "do-not-persist" not in json.dumps(report)
    assert emit_event.call_count == 2
    assert [call.kwargs["event_type"] for call in emit_event.call_args_list] == [
        "integration.sql_profile.completed",
        "integration.pimcore.completed",
    ]
    assert all(
        call.kwargs["job_id"] == report["operation_id"]
        for call in emit_event.call_args_list
    )
    assert all(call.kwargs["severity"] == "info" for call in emit_event.call_args_list)


def test_pimcore_history_change_set_includes_object_identity_and_stage_timings(monkeypatch):
    captured = {}
    monkeypatch.setattr(
        web_data,
        "record_history",
        lambda **kwargs: captured.update(kwargs) or kwargs,
    )
    report = {
        "operation_type": "manual_update",
        "username": "alice",
        "status": "completed",
        "total_ms": 91,
        "values": {"EAN": "5904804578169"},
        "events": [
            {"stage": "update", "stage_elapsed_ms": 31},
            {"stage": "verify", "stage_elapsed_ms": 17},
        ],
        "result": {
            "object": {"id": 91, "path": "/Produkty/ABC"},
            "change_set": {
                "kind": "updated",
                "fields": [
                    {"key": "SKU", "label": "SKU", "before": "OLD", "after": "NEW"}
                ],
            },
        },
    }

    web_data._persist_pimcore_operation(report)

    pimcore = captured["details"]["change_set"]["pimcore"]
    assert pimcore == {
        "kind": "updated",
        "fields": [
            {"key": "SKU", "label": "SKU", "before": "OLD", "after": "NEW"}
        ],
        "object_id": "91",
        "object_path": "/Produkty/ABC",
        "total_ms": 91,
        "send_ms": 31,
        "verification_ms": 17,
    }
def test_integration_results_use_trusted_requiredness_redact_errors_and_bound_warnings():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "REQUIRED_STOCK",
                    "label": "Required stock",
                    "pimcore_field": "requiredStock",
                    "type": "input",
                    "parser": "text",
                    "required": True,
                    "value_template": "SQL",
                    "sql_query": "SELECT 1",
                    "sql_profile_id": "required-profile",
                },
                {
                    "source": "OPTIONAL_STOCK",
                    "label": "Optional stock",
                    "pimcore_field": "optionalStock",
                    "type": "input",
                    "parser": "text",
                    "required": False,
                    "value_template": "SQL",
                    "sql_query": "SELECT 1",
                    "sql_profile_id": "optional-profile",
                },
            ],
        }
    )
    cfg["sql_profiles"] = [
        {"id": "required-profile", "password": "q7", "enabled": True},
        {"id": "optional-profile", "password": "server-optional-secret", "enabled": True},
    ]
    browser_results = {
        "sql_profiles": [
            {
                "profile_id": "required-profile",
                "source": "REQUIRED_STOCK",
                "status": "error",
                "required": False,
                "elapsed_ms": 8,
                "warning_codes": [f"warning-{index}" for index in range(40)],
                "error": (
                    "unlabeled q7; secret=browser secret with spaces; "
                    "password=browser-password; token=browser-token; "
                    "access_token=browser access token; "
                    "refresh_token='browser refresh token'; "
                    "client_secret=browser client secret; "
                    "db_password=browser database password; "
                    "x_api_key=browser compound api key; query_id=trace-42; "
                    + ("x" * 5000)
                ),
            },
            {
                "profile_id": "optional-profile",
                "source": "OPTIONAL_STOCK",
                "status": "error",
                "required": True,
                "elapsed_ms": 9,
                "warning_codes": [],
                "error": (
                    "server-optional-secret authorization=Bearer browser-auth "
                    "cookie=session=browser-cookie"
                ),
            },
        ]
    }
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "create_product",
            return_value={
                "created": True,
                "duplicate": False,
                "object": {"id": 91},
                "change_set": {"kind": "created", "fields": []},
            },
        ),
        patch.object(web_data, "record_history", return_value={}) as record_history,
        patch.object(web_data, "_persist_pimcore_submission"),
        patch.object(web_data, "emit_event") as emit_event,
    ):
        result = web_data.create_pimcore_product(
            {"EAN": "5904804578169"}, "operator", browser_results
        )

    profiles = result["change_set"]["integrations"]["sql_profiles"]
    assert [item["required"] for item in profiles] == [True, False]
    assert len(profiles[0]["warning_codes"]) == web_data.SQL_PROFILE_WARNING_CODES_MAX == 20
    assert len(profiles[0]["error"].encode("utf-8")) <= web_data.SQL_PROFILE_ERROR_MAX_BYTES
    serialized_result = json.dumps(result)
    history_details = record_history.call_args.kwargs["details"]
    serialized_history = json.dumps(history_details)
    serialized_events = json.dumps([call.kwargs["details"] for call in emit_event.call_args_list])
    for secret in (
        "q7",
        "server-optional-secret",
        "browser secret with spaces",
        "browser-password",
        "browser-token",
        "browser access token",
        "browser refresh token",
        "browser client secret",
        "browser database password",
        "browser compound api key",
        "browser-auth",
        "browser-cookie",
    ):
        assert secret not in serialized_result
        assert secret not in serialized_history
        assert secret not in serialized_events
    assert "query_id=trace-42" in profiles[0]["error"]
    assert [call.kwargs["severity"] for call in emit_event.call_args_list[:2]] == [
        "error",
        "info",
    ]


def test_pimcore_history_persists_common_change_set_with_nested_pimcore_diff():
    pimcore_change = {
        "kind": "updated",
        "fields": [{"key": "SKU", "label": "SKU", "before": "OLD", "after": "NEW"}],
    }
    integrations = {"sql_profiles": []}
    report = {
        "operation_type": "manual_update",
        "username": "operator",
        "values": {"EAN": "5904804578169"},
        "status": "completed",
        "result": {
            "object": {"id": 91},
            "change_set": {**pimcore_change, "integrations": integrations},
        },
        "integration_results": integrations,
    }
    with patch.object(web_data, "record_history", return_value={}) as record_history:
        web_data._persist_pimcore_operation(report)

    details = record_history.call_args.kwargs["details"]
    assert details["pimcore_operation"]["result"]["change_set"]["fields"] == pimcore_change[
        "fields"
    ]
    assert details["change_set"] == {
        "kind": "updated",
        "fields": [],
        "files": [],
        "integrations": integrations,
        "pimcore": {
            **pimcore_change,
            "integrations": integrations,
            "object_id": "91",
        },
    }


def test_create_adapter_persists_detailed_sqlite_submission_when_store_active():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    store = Mock()

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "_active_sqlite_store", return_value=store),
        patch.object(
            web_data,
            "create_product",
            return_value={
                "created": True,
                "duplicate": False,
                "object": {"id": 91},
                "payload": {"className": "Product"},
            },
        ),
        patch.object(web_data, "_persist_pimcore_operation"),
    ):
        web_data.create_pimcore_product({"EAN": "5904804578169"}, "operator")

    submitted = store.append_pimcore_submission.call_args.args[0]
    assert submitted["operation_type"] == "manual_create"
    assert submitted["username"] == "operator"
    assert submitted["values"]["EAN"] == "5904804578169"
    assert submitted["payload"]["className"] == "Product"


def test_create_adapter_persists_sqlite_submission_in_legacy_mode(tmp_path):
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    database_path = tmp_path / "pimcore-audit.sqlite"

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data.data_store,
            "get_active_store",
            return_value=web_data.data_store.LegacyDataStore(),
        ),
        patch.object(
            web_data.storage_settings,
            "resolve_sqlite_path",
            return_value=str(database_path),
        ),
        patch.object(
            web_data,
            "create_product",
            return_value={
                "created": True,
                "duplicate": False,
                "object": {"id": 91, "path": "/Produkty/5904"},
                "payload": {"className": "Product"},
            },
        ),
        patch.object(web_data, "_persist_pimcore_operation"),
    ):
        web_data.create_pimcore_product({"EAN": "5904804578169"}, "operator")

    rows = SqliteStore(str(database_path)).query_pimcore_submissions(
        user="operator",
        query="5904804578169",
        limit=10,
    )
    assert len(rows) == 1
    assert rows[0]["operation_type"] == "manual_create"
    assert rows[0]["payload"]["className"] == "Product"


def test_test_create_persists_sqlite_submission_when_operation_finishes():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    store = Mock()

    class ImmediateRegistry:
        def start(self, *, persist, **kwargs):
            report = {
                "operation_id": "op-test",
                "operation_type": kwargs["operation_type"],
                "username": kwargs["username"],
                "values": kwargs["values"],
                "status": "completed",
                "started_at": 1.0,
                "finished_at": 2.0,
                "events": [],
                "result": {"object_id": 77, "payload": {"className": "Product"}},
            }
            persist(report)
            return {"operation_id": "op-test", "status": "queued"}

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "_PIMCORE_OPERATIONS", ImmediateRegistry()),
        patch.object(web_data, "_active_sqlite_store", return_value=store),
        patch.object(web_data, "_persist_pimcore_operation") as persist_operation,
    ):
        result = web_data.start_pimcore_test_create(
            {"EAN": "5904804578169"},
            "keep",
            "operator",
        )

    assert result == {"operation_id": "op-test", "status": "queued"}
    persist_operation.assert_called_once()
    submitted = store.append_pimcore_submission.call_args.args[0]
    assert submitted["operation_type"] == "test"
    assert submitted["username"] == "operator"


def test_edit_adapter_persists_loaded_product_to_sqlite(tmp_path):
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update({"enabled": True, "setup_complete": True})
    database_path = tmp_path / "pimcore-audit.sqlite"
    loaded = {
        "object": {"id": 91, "path": "/Produkty/5904"},
        "marker": "100",
        "values": {"EAN": "5904804578169"},
    }

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data.storage_settings,
            "resolve_sqlite_path",
            return_value=str(database_path),
        ),
        patch.object(web_data, "fetch_product_for_edit", return_value=loaded),
        patch.object(web_data, "_persist_pimcore_operation"),
    ):
        result = web_data.get_pimcore_product_for_edit(91, "operator")

    assert result["object"]["id"] == 91
    rows = SqliteStore(str(database_path)).query_pimcore_submissions(
        operation_type="manual_load",
        user="operator",
        query="5904804578169",
        limit=10,
    )
    assert len(rows) == 1
    assert rows[0]["status"] == "completed"
    assert rows[0]["values"]["EAN"] == "5904804578169"


def test_runtime_edit_routes_allow_logged_in_user():
    client = TestClient(web_app.app)
    loaded = {
        "object": {"id": 91},
        "marker": "100",
        "values": {"EAN": "5904804578169"},
        "form_schema": [],
    }
    updated = {"object": {"id": 91}, "values": {"EAN": "5904804578169"}}
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(web_app, "get_pimcore_product_for_edit", return_value=loaded),
        patch.object(web_app, "update_pimcore_product", return_value=updated),
    ):
        get_response = client.get("/api/pimcore/products/91")
        put_response = client.put(
            "/api/pimcore/products/91",
            json={"marker": "100", "values": {"EAN": "5904804578169"}},
        )

    assert get_response.json() == loaded
    assert put_response.json() == updated


def test_runtime_edit_route_ignores_browser_supplied_integration_results():
    client = TestClient(web_app.app)
    integrations = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "success",
                "elapsed_ms": 8,
                "warning_codes": [],
                "error": "",
            }
        ]
    }
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(
            web_app,
            "update_pimcore_product",
            return_value={"object": {"id": 91}},
        ) as update,
    ):
        response = client.put(
            "/api/pimcore/products/91",
            json={
                "marker": "100",
                "values": {"EAN": "5904804578169"},
                "integration_results": integrations,
            },
        )

    assert response.status_code == 200
    update.assert_called_once_with(
        91,
        "100",
        {"EAN": "5904804578169"},
        "operator",
    )


def test_rendered_integration_context_is_bound_consumed_once_and_forwarded(tmp_path):
    client = TestClient(web_app.app)
    store = SqliteStore(str(tmp_path / "context.sqlite"))
    integrations = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "success",
                "elapsed_ms": 8,
                "warning_codes": [],
                "error": "",
            }
        ]
    }
    rendered = {
        "values": {"STOCK": "12"},
        "warnings": [],
        "integrations": integrations,
    }
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(web_app, "observability_store", return_value=store),
        patch.object(web_app, "render_saved_pimcore_templates", return_value=rendered),
        patch.object(
            web_app,
            "update_pimcore_product",
            return_value={"object": {"id": 91}},
        ) as update,
    ):
        render_response = client.post(
            "/api/pimcore/render-templates",
            json={
                "product_values": {"ean": "5904804578169"},
                "values": {"EAN": "5904804578169"},
                "targets": ["STOCK"],
                "mode": "edit",
                "object_id": 91,
            },
        )
        context_id = render_response.json()["integration_context_id"]
        first = client.put(
            "/api/pimcore/products/91",
            json={
                "marker": "100",
                "values": {"EAN": "5904804578169"},
                "integration_context_id": context_id,
                "integration_results": {"sql_profiles": [{"status": "fake-critical"}]},
            },
        )
        replay = client.put(
            "/api/pimcore/products/91",
            json={
                "marker": "100",
                "values": {"EAN": "5904804578169"},
                "integration_context_id": context_id,
            },
        )

    assert render_response.status_code == 200
    assert first.status_code == 200
    assert replay.status_code == 200
    assert update.call_args_list[0].args == (
        91,
        "100",
        {"EAN": "5904804578169"},
        "operator",
        integrations,
    )
    assert update.call_args_list[1].args == (
        91,
        "100",
        {"EAN": "5904804578169"},
        "operator",
    )


def test_runtime_edit_conflict_returns_409():
    client = TestClient(web_app.app)
    error = PimcoreConflictError("Obiekt zostal zmieniony.", 91, "100", "101")
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(web_app, "update_pimcore_product", side_effect=error),
    ):
        response = client.put(
            "/api/pimcore/products/91",
            json={"marker": "100", "values": {"EAN": "5904804578169"}},
        )
    assert response.status_code == 409
    assert response.json()["detail"]["current_marker"] == "101"


def test_admin_can_preview_unsaved_pimcore_template():
    client = TestClient(web_app.app)
    payload = {
        "mappings": [],
        "target_source": "TITLE",
        "product_values": {},
        "values": {},
    }
    expected = {"values": {"TITLE": "VIVO"}, "warnings": []}
    with (
        patch.object(web_app, "_require_admin", return_value="admin"),
        patch.object(
            web_app,
            "preview_pimcore_template",
            return_value=expected,
            create=True,
        ),
    ):
        response = client.post(
            "/api/settings/pimcore/template-preview",
            json=payload,
        )

    assert response.status_code == 200
    assert response.json() == expected


def test_template_preview_fills_missing_product_placeholders_from_saved_entry():
    payload = {
        "mappings": [
            {
                "source": "TITLE",
                "label": "Nazwa",
                "pimcore_field": "title",
                "type": "input",
                "parser": "text",
                "value_template": "{PRODUCT:name|keep} - {PRODUCT:type|keep}",
            }
        ],
        "target_source": "TITLE",
        "product_values": {},
        "values": {},
    }
    records = {
        web_data.ENTRY_RECORDS_KEY: [
            {
                web_data.NAME_HEADER: "Vivo",
                web_data.TYPE_HEADER: "Komoda",
                web_data.MODEL_HEADER: "M1",
                web_data.COLOR1_HEADER: "bialy",
                web_data.EAN_HEADER: "5904804578169",
            }
        ]
    }

    with patch.object(web_data, "prepare_excel_lists", return_value=records):
        result = web_data.preview_pimcore_template(payload)

    assert result["values"]["TITLE"] == "Vivo - Komoda"


def test_admin_test_sample_route_returns_fresh_editable_values():
    client = TestClient(web_app.app)
    expected = {
        "form_schema": [{"source": "EAN"}],
        "values": {"EAN": "5904804578169"},
        "warnings": [],
    }
    with (
        patch.object(web_app, "_require_admin", return_value="admin"),
        patch.object(
            web_app,
            "pimcore_test_sample",
            return_value=expected,
            create=True,
        ),
    ):
        response = client.post("/api/settings/pimcore/test-sample")

    assert response.status_code == 200
    assert response.json() == expected


def test_test_sample_renders_product_placeholders_from_saved_entry_when_available():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "TITLE",
                    "label": "Title",
                    "pimcore_field": "title",
                    "type": "input",
                    "parser": "text",
                    "value_template": "{PRODUCT:name|keep} - {PRODUCT:type|keep}",
                }
            ],
        }
    )
    records = {
        web_data.ENTRY_RECORDS_KEY: [
            {
                web_data.NAME_HEADER: "Vivo",
                web_data.TYPE_HEADER: "Komoda",
                web_data.MODEL_HEADER: "M1",
                web_data.COLOR1_HEADER: "bialy",
                web_data.EAN_HEADER: "5904804578169",
            }
        ]
    }

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "prepare_excel_lists", return_value=records),
    ):
        sample = web_data.pimcore_test_sample()

    assert sample["values"]["TITLE"] == "Vivo - Komoda"


def test_test_sample_is_available_when_complete_integration_is_temporarily_disabled():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": False,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "SKU",
                    "label": "SKU",
                    "pimcore_field": "sku",
                    "type": "input",
                    "parser": "text",
                }
            ],
        }
    )

    with patch.object(web_data.config, "CONFIG", cfg):
        sample = web_data.pimcore_test_sample()

    assert sample["values"]["SKU"].startswith("TEST_SKU_")


def test_logged_in_user_can_render_only_saved_templates():
    client = TestClient(web_app.app)
    expected = {"values": {"TITLE": "VIVO"}, "warnings": []}
    with (
        patch.object(web_app, "_require_user", return_value="operator"),
        patch.object(
            web_app,
            "render_saved_pimcore_templates",
            return_value=expected,
            create=True,
        ) as render,
    ):
        response = client.post(
            "/api/pimcore/render-templates",
            json={
                "product_values": {"name": "Vivo"},
                "values": {},
                "targets": ["TITLE"],
            },
        )

    assert response.status_code == 200
    assert response.json() == expected
    render.assert_called_once_with({"name": "Vivo"}, {}, ["TITLE"], "create")


def test_render_saved_pimcore_templates_auto_applies_sql_only_when_empty():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stan",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "value_template": "SQL",
                    "sql_query": "SELECT stock FROM product WHERE ean = {ean}",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [
        {
            "id": "stock",
            "label": "Stock",
            "type": "mysql",
            "host": "mysql.local",
            "database": "catalog",
            "user": "reader",
            "password": "secret",
            "enabled": True,
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "execute_sql_value_query",
            return_value=web_data.SqlValueResult("12", []),
        ),
    ):
        empty = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"},
            {"STOCK": ""},
            ["STOCK"],
        )
        manual = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"},
            {"STOCK": "manual"},
            ["STOCK"],
        )

    assert empty["values"]["STOCK"] == "12"
    assert empty["calculated_values"]["STOCK"] == "12"
    assert empty["changed"]["STOCK"] is False
    assert manual["values"]["STOCK"] == "manual"
    assert manual["calculated_values"]["STOCK"] == "12"
    assert manual["changed"]["STOCK"] is True


def test_render_saved_pimcore_templates_reports_each_sql_profile_integration():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stan",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "value_template": "SQL",
                    "sql_query": "SELECT stock FROM product WHERE ean = {ean}",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [
        {
            "id": "stock",
            "label": "Stock",
            "type": "mysql",
            "host": "mysql.local",
            "database": "catalog",
            "user": "reader",
            "password": "secret",
            "enabled": True,
        }
    ]
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "execute_sql_value_query",
            return_value=web_data.SqlValueResult(
                "12",
                [{"code": "multiple_rows", "message": "Used first row"}]
                + [
                    {"code": f"warning-{index}", "message": "warning"}
                    for index in range(40)
                ],
            ),
        ),
    ):
        result = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"},
            {"STOCK": ""},
            ["STOCK"],
        )

    integration = result["integrations"]["sql_profiles"][0]
    assert integration["profile_id"] == "stock"
    assert integration["source"] == "STOCK"
    assert integration["status"] == "warning"
    assert integration["warning_codes"][0] == "multiple_rows"
    assert len(integration["warning_codes"]) == web_data.SQL_PROFILE_WARNING_CODES_MAX
    assert integration["error"] == ""
    assert integration["elapsed_ms"] >= 0


def test_render_sql_profile_error_redacts_assignments_and_known_profile_secret():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stock",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "required": True,
                    "value_template": "SQL",
                    "sql_query": "SELECT stock",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [{"id": "stock", "password": "q7", "enabled": True}]
    raw_error = (
        "unlabeled q7; PWD=driver secret with spaces; "
        "password='inline secret with spaces'; pass=pass secret with spaces; "
        "secret=plain secret with spaces; token=token secret with spaces; "
        "authorization=Bearer auth secret with spaces; "
        "api_key=underscore api secret; api-key=hyphen api secret; "
        "cookie=session cookie secret; access_token=render access token; "
        "refresh_token='render refresh token'; "
        "client_secret=render client secret; db_password=render database password; "
        "x_api_key=render compound api key; query_id=trace-42\nsafe diagnostic"
    )
    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(web_data, "execute_sql_value_query", side_effect=RuntimeError(raw_error)),
    ):
        result = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"}, {"STOCK": ""}, ["STOCK"]
        )

    serialized = json.dumps(result)
    for secret in (
        "q7",
        "driver secret with spaces",
        "inline secret with spaces",
        "pass secret with spaces",
        "plain secret with spaces",
        "token secret with spaces",
        "auth secret with spaces",
        "underscore api secret",
        "hyphen api secret",
        "session cookie secret",
        "render access token",
        "render refresh token",
        "render client secret",
        "render database password",
        "render compound api key",
    ):
        assert secret not in serialized
    assert "[REDACTED]" in serialized
    integration = result["integrations"]["sql_profiles"][0]
    assert integration["required"] is True
    assert "query_id=trace-42" in integration["error"]


def test_sql_warning_codes_caps_before_iterating_warning_collection():
    class SliceOnlyWarnings(list):
        def __iter__(self):
            raise AssertionError("warning collection was traversed before capping")

    warnings = SliceOnlyWarnings(
        [{"code": f"warning-{index}"} for index in range(1000)]
    )

    assert web_data._sql_warning_codes(warnings) == [
        f"warning-{index}" for index in range(web_data.SQL_PROFILE_WARNING_CODES_MAX)
    ]


def test_safe_integration_results_caps_browser_warning_codes_before_iterating():
    class SliceOnlyCodes(list):
        def __iter__(self):
            raise AssertionError("warning codes were traversed before capping")

    settings = {
        "field_mappings": [
            {
                "source": "STOCK",
                "required": False,
                "value_template": "SQL",
                "sql_profile_id": "stock",
            }
        ]
    }
    submitted = {
        "sql_profiles": [
            {
                "profile_id": "stock",
                "source": "STOCK",
                "status": "warning",
                "warning_codes": SliceOnlyCodes(
                    [f"warning-{index}" for index in range(1000)]
                ),
            }
        ]
    }

    result = web_data._safe_pimcore_integration_results(submitted, settings)

    assert result["sql_profiles"][0]["warning_codes"] == [
        f"warning-{index}" for index in range(web_data.SQL_PROFILE_WARNING_CODES_MAX)
    ]


def test_render_saved_pimcore_templates_uses_sql_as_template_source():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK_LABEL",
                    "label": "Stan",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "value_template": "Stan: {SQL|number:0}",
                    "sql_query": "SELECT stock FROM product WHERE ean = {ean}",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [
        {
            "id": "stock",
            "label": "Stock",
            "type": "mysql",
            "host": "mysql.local",
            "database": "catalog",
            "user": "reader",
            "password": "secret",
            "enabled": True,
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "execute_sql_value_query",
            return_value=web_data.SqlValueResult("12.4", []),
        ) as execute_sql,
    ):
        result = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"},
            {"STOCK_LABEL": ""},
            ["STOCK_LABEL"],
        )

    assert result["values"]["STOCK_LABEL"] == "Stan: 12"
    assert result["calculated_values"]["STOCK_LABEL"] == "Stan: 12"
    execute_sql.assert_called_once()


def test_render_saved_pimcore_templates_marks_sql_template_source_difference_for_edit():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "SKU",
                    "label": "SKU",
                    "pimcore_field": "sku",
                    "type": "input",
                    "parser": "text",
                    "value_template": "{SQL|keep}",
                    "sql_query": "SELECT TOP 1 sku FROM product WHERE ean = '{PRODUCT:ean|keep}'",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [
        {
            "id": "stock",
            "label": "Stock",
            "type": "mssql",
            "host": "sql.local",
            "database": "catalog",
            "user": "reader",
            "password": "secret",
            "enabled": True,
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "execute_sql_value_query",
            return_value=web_data.SqlValueResult("SKU-NEW", []),
        ),
    ):
        result = web_data.render_saved_pimcore_templates(
            {"ean": "5907763645590"},
            {"SKU": "SKU-OLD"},
            ["SKU"],
            mode="edit",
        )

    assert result["values"]["SKU"] == "SKU-NEW"
    assert result["calculated_values"]["SKU"] == "SKU-NEW"
    assert result["changed"]["SKU"] is True


def test_render_saved_pimcore_templates_for_edit_does_not_auto_apply_sql():
    cfg = json.loads(json.dumps(web_data.config.DEFAULT_CONFIG))
    cfg["pimcore"].update(
        {
            "enabled": True,
            "setup_complete": True,
            "field_mappings": [
                {
                    "source": "STOCK",
                    "label": "Stan",
                    "pimcore_field": "stockText",
                    "type": "input",
                    "parser": "text",
                    "value_template": "SQL",
                    "sql_query": "SELECT stock FROM product WHERE ean = {ean}",
                    "sql_profile_id": "stock",
                }
            ],
        }
    )
    cfg["sql_profiles"] = [
        {
            "id": "stock",
            "label": "Stock",
            "type": "mysql",
            "host": "mysql.local",
            "database": "catalog",
            "user": "reader",
            "password": "secret",
            "enabled": True,
        }
    ]

    with (
        patch.object(web_data.config, "CONFIG", cfg),
        patch.object(
            web_data,
            "execute_sql_value_query",
            return_value=web_data.SqlValueResult("12", []),
        ),
    ):
        result = web_data.render_saved_pimcore_templates(
            {"ean": "5901234567890"},
            {"STOCK": "existing"},
            ["STOCK"],
            mode="edit",
        )

    assert result["values"]["STOCK"] == "existing"
    assert result["calculated_values"]["STOCK"] == "12"
    assert result["changed"]["STOCK"] is True
